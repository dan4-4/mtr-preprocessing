from __future__ import annotations

import re
from difflib import SequenceMatcher
from pathlib import Path

import openpyxl


def _normalize_text(text: str) -> str:
    text = text.lower().strip()
    text = re.sub(r"\s+", " ", text)
    return text


def _token_set(text: str) -> set[str]:
    return {t for t in re.findall(r"[a-zа-яё0-9]+", _normalize_text(text)) if len(t) > 2}


class MaterialCatalog:
    """
    Catalog from «Пример отчета мониторинга цен.xlsx».
    Columns: F (source name), P/V/AB (normalized supplier names).
    """

    CATALOG_NAME_COLUMNS = ("Наименование 1", "Наименование 2", "Наименование 3")
    SOURCE_NAME_COLUMN = "Наименование"

    def __init__(self, catalog_path: str | Path):
        self._entries: list[dict] = []
        self._normalized_names: list[str] = []
        self._load(catalog_path)

    def _load(self, catalog_path: str | Path) -> None:
        wb = openpyxl.load_workbook(catalog_path, read_only=True, data_only=True)
        try:
            ws = wb[wb.sheetnames[0]]
            header = list(next(ws.iter_rows(min_row=1, values_only=True)))
            col_index = {str(h): i for i, h in enumerate(header) if h is not None}

            source_idx = col_index.get(self.SOURCE_NAME_COLUMN)
            catalog_idxs = [
                col_index[name]
                for name in self.CATALOG_NAME_COLUMNS
                if name in col_index
            ]

            for row in ws.iter_rows(min_row=2, values_only=True):
                source_name = (
                    str(row[source_idx]).strip()
                    if source_idx is not None and row[source_idx]
                    else ""
                )
                normalized_values: list[str] = []
                for idx in catalog_idxs:
                    val = row[idx]
                    if val:
                        normalized_values.append(str(val).strip())

                if not source_name and not normalized_values:
                    continue

                self._entries.append(
                    {"source_name": source_name, "normalized_names": normalized_values}
                )
                self._normalized_names.extend(normalized_values)
                if source_name:
                    self._normalized_names.append(source_name)
        finally:
            wb.close()

        seen: set[str] = set()
        unique: list[str] = []
        for name in self._normalized_names:
            key = _normalize_text(name)
            if key and key not in seen:
                seen.add(key)
                unique.append(name)
        self._normalized_names = unique

    def find_similar_source(self, name_raw: str, top_k: int = 3) -> list[str]:
        if not name_raw:
            return []
        query_norm = _normalize_text(name_raw)
        scored: list[tuple[float, list[str]]] = []
        for entry in self._entries:
            source = entry["source_name"]
            if not source:
                continue
            ratio = SequenceMatcher(None, query_norm, _normalize_text(source)).ratio()
            if ratio >= 0.5:
                scored.append((ratio, entry["normalized_names"]))

        scored.sort(key=lambda x: x[0], reverse=True)
        results: list[str] = []
        seen: set[str] = set()
        for _, names in scored:
            for name in names:
                key = _normalize_text(name)
                if key and key not in seen:
                    seen.add(key)
                    results.append(name)
            if len(results) >= top_k:
                break
        return results[:top_k]

    def search(self, query: str, top_k: int = 5, min_score: float = 0.35) -> list[str]:
        if not query or not self._normalized_names:
            return []

        query_norm = _normalize_text(query)
        query_tokens = _token_set(query)
        scored: list[tuple[float, str]] = []
        for name in self._normalized_names:
            name_norm = _normalize_text(name)
            ratio = SequenceMatcher(None, query_norm, name_norm).ratio()
            overlap = len(query_tokens & _token_set(name)) / max(len(query_tokens), 1)
            score = 0.6 * ratio + 0.4 * overlap
            if score >= min_score:
                scored.append((score, name))

        scored.sort(key=lambda x: x[0], reverse=True)
        results: list[str] = []
        seen: set[str] = set()
        for _, name in scored:
            key = _normalize_text(name)
            if key in seen:
                continue
            seen.add(key)
            results.append(name)
            if len(results) >= top_k:
                break
        return results

    def matches_for(self, name_raw: str, top_k: int = 5) -> list[str]:
        from_source = self.find_similar_source(name_raw, top_k=3)
        if from_source:
            return from_source[:top_k]
        return self.search(name_raw, top_k=top_k)

    def __len__(self) -> int:
        return len(self._normalized_names)
