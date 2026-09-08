from __future__ import annotations

import re
from pathlib import Path

import openpyxl

from preprocessing.constants import KSR_CODE_PATTERN, KSR_ROW_PATTERN


class KsrLookup:
    """Lookup KSR code → name from Выгрузка КСР.xlsx."""

    def __init__(self, reference_path: str | Path):
        self._by_code: dict[str, str] = {}
        self._load(reference_path)

    def _load(self, reference_path: str | Path) -> None:
        wb = openpyxl.load_workbook(reference_path, read_only=True, data_only=True)
        try:
            for sheet_name in wb.sheetnames:
                ws = wb[sheet_name]
                for row in ws.iter_rows(min_row=3, values_only=True):
                    if not row or row[0] is None:
                        continue
                    code = str(row[0]).strip()
                    if not KSR_ROW_PATTERN.match(code):
                        continue
                    name = str(row[1]).strip() if row[1] else ""
                    self._by_code[code] = name
        finally:
            wb.close()

    def lookup(self, code: str | None) -> str | None:
        if not code:
            return None
        return self._by_code.get(code)

    def __len__(self) -> int:
        return len(self._by_code)


def extract_ksr_code(justification: str | None) -> str | None:
    if not justification:
        return None
    match = KSR_CODE_PATTERN.search(str(justification))
    return match.group(0) if match else None
