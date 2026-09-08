from __future__ import annotations

from pathlib import Path
from typing import Any

import openpyxl

from preprocessing.constants import REQUIRED_SOURCE_FIELDS, SOURCE_COLUMNS
from preprocessing.ksr_lookup import KsrLookup, extract_ksr_code
from preprocessing.models import ResourceItem


def _coerce_float(value: Any) -> float | None:
    if value is None or value == "":
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def _coerce_int(value: Any) -> int | None:
    if value is None or value == "":
        return None
    try:
        return int(value)
    except (TypeError, ValueError):
        return None


def _get_cell(row: tuple, field: str) -> Any:
    idx = SOURCE_COLUMNS[field]
    return row[idx] if idx < len(row) else None


def parse_excel_to_resource_items(
    excel_path: str | Path,
    ksr_reference_path: str | Path,
    *,
    sheet_name: str | None = None,
    id_prefix: str = "raw",
) -> list[ResourceItem]:
    """
    Stage 0: parse source Excel into ResourceItem list.

    Expected columns (0-based): № КС, Шифр сметы, тип ресурса, _, Обоснование,
    Наименование, Ед. изм., Расход, Базовая цена, Стоимость.
    """
    excel_path = Path(excel_path)
    ksr_lookup = KsrLookup(ksr_reference_path)

    wb = openpyxl.load_workbook(excel_path, read_only=True, data_only=True)
    try:
        ws = wb[sheet_name] if sheet_name else wb[wb.sheetnames[0]]
        items: list[ResourceItem] = []

        for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=1):
            if not row or all(cell is None for cell in row):
                continue

            values = {field: _get_cell(row, field) for field in SOURCE_COLUMNS}
            name_raw = str(values["name_raw"]).strip() if values["name_raw"] else ""
            unit_raw = str(values["unit_raw"]).strip() if values["unit_raw"] else ""

            if not name_raw and not unit_raw:
                continue

            missing = [f for f in REQUIRED_SOURCE_FIELDS if not values.get(f)]
            if missing:
                raise ValueError(
                    f"Row {row_idx + 1}: missing required fields: {', '.join(missing)}"
                )

            justification = (
                str(values["justification_source"]).strip()
                if values["justification_source"]
                else None
            )
            ksr_code = extract_ksr_code(justification)
            ksr_name = ksr_lookup.lookup(ksr_code)

            items.append(
                ResourceItem(
                    item_id=f"{id_prefix}-{row_idx:04d}",
                    ks_number=_coerce_int(values["ks_number"]),
                    estimate_code=(
                        str(values["estimate_code"]).strip()
                        if values["estimate_code"]
                        else None
                    ),
                    resource_type=_coerce_int(values["resource_type"]),
                    justification_source=justification,
                    ksr_code=ksr_code,
                    ksr_name=ksr_name,
                    name_raw=name_raw,
                    unit_raw=unit_raw,
                    consumption=_coerce_float(values["consumption"]),
                    base_price=_coerce_float(values["base_price"]),
                    total_cost=_coerce_float(values["total_cost"]),
                )
            )
    finally:
        wb.close()

    return items
